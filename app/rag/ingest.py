"""把 HR 文档解析、切分、向量化并写入 Chroma。

用法：
    python -m app.rag.ingest
    python -m app.rag.ingest --docs-dir data/docs --no-reset
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
from pathlib import Path
from typing import Iterable

from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.config import PROJECT_ROOT, settings

SUPPORTED_SUFFIXES = {".md", ".txt", ".pdf", ".docx"}
COLLECTION_NAME = "hr_docs"

_enc = None
try:  # tiktoken 首次使用可能需要联网下载编码表，失败时退回正则估算。
    import tiktoken

    _enc = tiktoken.get_encoding("cl100k_base")
except Exception:  # noqa: BLE001
    _enc = None


def token_length(text: str) -> int:
    if _enc is not None:
        return len(_enc.encode(text))
    return len(re.findall(r"[一-鿿]|[A-Za-z0-9_]+", text))


def _read_pdf(path: Path) -> str:
    import pymupdf

    parts: list[str] = []
    with pymupdf.open(path) as doc:
        for page in doc:
            parts.append(page.get_text("text"))
    return "\n\n".join(parts)


def _read_docx(path: Path) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(path)
    lines: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower()
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            lines.append("#" * max(1, min(int(level), 6)) + " " + text)
        else:
            lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)


def read_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return _read_docx(path)
    return path.read_text(encoding="utf-8", errors="ignore")


def iter_files(docs_dir: Path) -> Iterable[Path]:
    if not docs_dir.exists():
        return []
    return sorted(
        p
        for p in docs_dir.rglob("*")
        if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES
    )


def split_by_headings(text: str, fallback_heading: str) -> list[tuple[str, str]]:
    """按 Markdown 标题切 section，保留标题行作为正文的一部分。"""
    sections: list[tuple[str, str]] = []
    heading = fallback_heading
    buffer: list[str] = []

    def flush() -> None:
        body = "\n".join(buffer).strip()
        if body:
            sections.append((heading, body))

    for line in text.splitlines():
        if re.match(r"^#{1,6}\s+\S", line):
            flush()
            heading = line.lstrip("#").strip()
            buffer = [line]
        else:
            buffer.append(line)
    flush()
    return sections


def build_chunks(docs_dir: Path) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=500,
        chunk_overlap=80,
        length_function=token_length,
        separators=["\n\n", "\n", "。", "！", "？", "；", "，", " ", ""],
    )

    chunks: list[Document] = []
    for path in iter_files(docs_dir):
        relative = path.relative_to(docs_dir).as_posix()
        text = read_file(path)
        sections = split_by_headings(text, path.stem)
        for section_index, (heading, body) in enumerate(sections):
            for chunk_index, piece in enumerate(splitter.split_text(body)):
                piece = piece.strip()
                if not piece:
                    continue
                chunk_id = f"{relative}::{section_index}::{chunk_index}"
                chunks.append(
                    Document(
                        page_content=piece,
                        metadata={
                            "chunk_id": chunk_id,
                            "source": relative,
                            "source_name": path.name,
                            "heading": heading,
                            "section_index": section_index,
                            "chunk_index": chunk_index,
                        },
                    )
                )
    return chunks


def build_vectorstore(chunks: list[Document], reset: bool = True) -> Chroma:
    if reset:
        shutil.rmtree(settings.chroma_path, ignore_errors=True)
    settings.chroma_path.mkdir(parents=True, exist_ok=True)

    embeddings = HuggingFaceEmbeddings(
        model_name=settings.embed_model,
        model_kwargs={"device": "cpu"},
        encode_kwargs={"normalize_embeddings": True, "batch_size": 8},
    )
    store = Chroma(
        collection_name=COLLECTION_NAME,
        persist_directory=str(settings.chroma_path),
        embedding_function=embeddings,
    )
    for start in range(0, len(chunks), 32):
        batch = chunks[start : start + 32]
        store.add_documents(
            batch,
            ids=[doc.metadata["chunk_id"] for doc in batch],
        )
        print(f"  已写入 {min(start + 32, len(chunks))}/{len(chunks)}", flush=True)
    return store


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--docs-dir",
        type=Path,
        default=settings.docs_path,
        help="HR 文档目录，默认 data/docs",
    )
    parser.add_argument("--no-reset", action="store_true", help="不清空已有 Chroma 集合")
    args = parser.parse_args()

    docs_dir = args.docs_dir
    if not docs_dir.is_absolute():
        docs_dir = PROJECT_ROOT / docs_dir

    files = list(iter_files(docs_dir))
    print(f"文档目录：{docs_dir}", flush=True)
    print(f"发现 {len(files)} 个文件", flush=True)
    if not files:
        print("没有可解析的文档，请先放入 .md/.txt/.pdf/.docx 文件。", flush=True)
        return 1

    chunks = build_chunks(docs_dir)
    if not chunks:
        print("没有生成任何 chunk。", flush=True)
        return 1

    settings.chunks_path.parent.mkdir(parents=True, exist_ok=True)
    with settings.chunks_path.open("w", encoding="utf-8") as f:
        for doc in chunks:
            f.write(
                json.dumps(
                    {"content": doc.page_content, "metadata": doc.metadata},
                    ensure_ascii=False,
                )
                + "\n"
            )
    print(f"已生成 {len(chunks)} 个 chunk -> {settings.chunks_path}", flush=True)

    build_vectorstore(chunks, reset=not args.no_reset)
    print(f"向量库已写入 -> {settings.chroma_path}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
